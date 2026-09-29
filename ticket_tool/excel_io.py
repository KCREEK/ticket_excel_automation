"""Excel helpers - reproduces the structure of the 'Ticketing Sheet' template.

Template (Demo_sheets.xlsx)
---------------------------
Sheet name : Ticketing Sheet
Row 1      : ticket_no | ticket_priority | ticket_status | ticket_sub |
             action_taken | time_taken       (Calibri 11, bold, red, centered,
             top-aligned, wrapped, thin borders, row height 15)
Rows 2+    : Calibri 11, centered both ways, wrapped, thin borders,
             row height 30 (template is pre-formatted down to row 760)
Widths     : A 16.55, B 12.33, C 23.33, D 50, E 50.55, F 10

Existing rows are never modified: new rows are written below the last row
that has data. Duplicate ticket IDs are never dropped.
"""
from __future__ import annotations

import math
import os
import re
import shutil
import tempfile
from datetime import datetime
from typing import Dict, List, Optional, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .parser import TicketRow

DEFAULT_SHEET = "Ticketing Sheet"
HEADERS = ["ticket_no", "ticket_priority", "ticket_status", "ticket_sub",
           "action_taken", "time_taken"]
COLUMN_WIDTHS = [16.5546875, 12.33203125, 23.33203125, 50.0, 50.5546875, 10.0]
WHITE_FILL_COLS = (1, 4, 5)          # A, D, E carry a white fill in the template
PREFORMAT_ROWS = 760
BODY_ROW_HEIGHT = 30.0
MAX_ROW_HEIGHT = 409.0               # Excel's hard limit
MAX_CELL_CHARS = 32767               # Excel's hard limit

_THIN = Side(style="thin")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


class ExcelWriteError(Exception):
    """Raised with a user-friendly message."""


# ---------------------------------------------------------------- styling
def _style_header(cell) -> None:
    cell.font = Font(name="Calibri", size=11, bold=True, color="FFFF0000")
    cell.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
    cell.border = _BORDER


def _style_body(cell) -> None:
    cell.font = Font(name="Calibri", size=11)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = _BORDER
    if cell.column in WHITE_FILL_COLS:
        cell.fill = PatternFill("solid", fgColor="FFFFFFFF")


def create_new_workbook(path: str) -> None:
    """Create a fresh workbook with the same structure as the template."""
    wb = Workbook()
    ws = wb.active
    ws.title = DEFAULT_SHEET
    for i, (name, width) in enumerate(zip(HEADERS, COLUMN_WIDTHS), start=1):
        _style_header(ws.cell(row=1, column=i, value=name))
        ws.column_dimensions[chr(64 + i)].width = width
    ws.row_dimensions[1].height = 15.0
    for r in range(2, PREFORMAT_ROWS + 1):
        for c in range(1, len(HEADERS) + 1):
            _style_body(ws.cell(row=r, column=c))
        ws.row_dimensions[r].height = BODY_ROW_HEIGHT
    wb.save(path)


# ---------------------------------------------------------------- reading
def _open(path: str, **kw):
    keep_vba = path.lower().endswith(".xlsm")
    try:
        return load_workbook(path, keep_vba=keep_vba, **kw)
    except PermissionError:
        raise ExcelWriteError(f"Cannot open '{os.path.basename(path)}'. "
                              "Close it in Excel and try again.")
    except Exception as exc:                       # corrupt / not an xlsx
        raise ExcelWriteError(f"Could not open the Excel file:\n{exc}")


def list_sheets(path: str) -> List[str]:
    wb = _open(path, read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def detect_columns(ws) -> List[int]:
    """Map our six fields to column numbers, by header name (row 1),
    falling back to the template's positions A-F."""
    found: Dict[str, int] = {}
    for c in range(1, min(ws.max_column, 30) + 1):
        v = ws.cell(row=1, column=c).value
        if isinstance(v, str):
            found.setdefault(v.strip().lower(), c)
    return [found.get(h, i + 1) for i, h in enumerate(HEADERS)]


def headers_match(ws) -> bool:
    return all(str(ws.cell(row=1, column=c).value or "").strip().lower() == h
               for c, h in zip(detect_columns(ws), HEADERS))


def _last_data_row(ws, cols: Sequence[int]) -> int:
    for r in range(ws.max_row, 1, -1):
        if any(ws.cell(row=r, column=c).value not in (None, "") for c in cols):
            return r
    return 1


def inspect_sheet(path: str, sheet: Optional[str] = None) -> Dict:
    """Small summary shown in the GUI so you can sanity-check the target."""
    wb = _open(path)
    try:
        ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.worksheets[0]
        cols = detect_columns(ws)
        last = _last_data_row(ws, cols)
        return {"sheet": ws.title, "last_row": last, "next_row": last + 1,
                "existing_rows": max(0, last - 1), "headers_ok": headers_match(ws)}
    finally:
        wb.close()


def _norm(v) -> str:
    return "" if v is None else str(v).replace("\r\n", "\n").strip()


def count_existing_matches(path: str, rows: Sequence[TicketRow],
                           sheet: Optional[str] = None) -> int:
    """How many of `rows` are already in the sheet, identical in every text
    column. Informational only - used to warn against running the same
    notes file twice. Nothing is ever skipped automatically."""
    wb = _open(path)
    try:
        ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.worksheets[0]
        cols = detect_columns(ws)[:5]
        existing = set()
        for r in range(2, ws.max_row + 1):
            existing.add(tuple(_norm(ws.cell(row=r, column=c).value) for c in cols))
        existing.discard(("",) * 5)
        return sum(1 for r in rows
                   if tuple(_norm(x) for x in r.as_list()[:5]) in existing)
    finally:
        wb.close()


# ---------------------------------------------------------------- writing
def _clean(v: str) -> str:
    v = _ILLEGAL.sub("", v or "")
    return v[:MAX_CELL_CHARS]


def _time_value(v: str):
    v = (v or "").strip()
    if re.fullmatch(r"\d+", v):
        return int(v)
    if re.fullmatch(r"\d+\.\d+", v):
        return float(v)
    return v


def _estimate_height(values: Sequence[str], widths: Sequence[float]) -> float:
    lines = 1
    for text, width in zip(values, widths):
        per_line = max(1, int(width * 1.1))
        n = sum(max(1, math.ceil(len(seg) / per_line))
                for seg in str(text).split("\n"))
        lines = max(lines, n)
    return min(MAX_ROW_HEIGHT, max(BODY_ROW_HEIGHT, lines * 15.0))


def make_backup(path: str) -> str:
    folder = os.path.join(os.path.dirname(os.path.abspath(path)), "backups")
    os.makedirs(folder, exist_ok=True)
    stem, ext = os.path.splitext(os.path.basename(path))
    dest = os.path.join(folder, f"{stem}_backup_{datetime.now():%Y%m%d_%H%M%S}{ext}")
    shutil.copy2(path, dest)
    return dest


def append_rows(path: str, rows: Sequence[TicketRow], sheet: Optional[str] = None,
                backup: bool = True, autofit_height: bool = True) -> Dict:
    """Append `rows` below the last used row. Every row is written - repeated
    ticket IDs are kept. Existing cells are not touched."""
    if not rows:
        raise ExcelWriteError("There are no rows to write.")
    wb = _open(path)
    ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.worksheets[0]
    cols = detect_columns(ws)
    start = _last_data_row(ws, cols) + 1
    widths = [ws.column_dimensions[chr(64 + c)].width or 13.0 if c <= 26 else 13.0
              for c in cols]

    for offset, row in enumerate(rows):
        r = start + offset
        values = [_clean(row.ticket_id), _clean(row.priority), _clean(row.status),
                  _clean(row.title), _clean(row.description),
                  _time_value(row.time_taken)]
        for col, val in zip(cols, values):
            cell = ws.cell(row=r, column=col)
            cell.value = val
            if isinstance(val, str) and val.startswith("="):
                cell.data_type = "s"               # never let notes become formulas
            if cell.border is None or cell.border.left.style is None:
                _style_body(cell)                  # row beyond the pre-formatted area
        current = ws.row_dimensions[r].height or BODY_ROW_HEIGHT
        target = (_estimate_height([str(v) for v in values], widths)
                  if autofit_height else BODY_ROW_HEIGHT)
        ws.row_dimensions[r].height = max(current, target) if autofit_height else current

    backup_path = make_backup(path) if backup else None
    fd, tmp = tempfile.mkstemp(suffix=os.path.splitext(path)[1],
                               dir=os.path.dirname(os.path.abspath(path)))
    os.close(fd)
    try:
        wb.save(tmp)
        os.replace(tmp, path)
    except PermissionError:
        raise ExcelWriteError(f"'{os.path.basename(path)}' is open in Excel. "
                              "Close it and press the button again.")
    finally:
        wb.close()
        if os.path.exists(tmp):
            os.remove(tmp)
    return {"sheet": ws.title, "first_row": start, "last_row": start + len(rows) - 1,
            "count": len(rows), "backup": backup_path}
