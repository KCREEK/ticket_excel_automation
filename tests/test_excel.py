import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import load_workbook  # noqa: E402

from ticket_tool import excel_io  # noqa: E402
from ticket_tool.parser import TicketRow, parse_file  # noqa: E402

SAMPLE = os.path.join(os.path.dirname(__file__), "..", "samples", "sample_notes.txt")


def test_new_workbook_matches_template(tmp_path):
    p = str(tmp_path / "new.xlsx")
    excel_io.create_new_workbook(p)
    ws = load_workbook(p).active
    assert ws.title == "Ticketing Sheet"
    assert [c.value for c in ws[1]] == excel_io.HEADERS
    assert ws["A1"].font.b and ws["A1"].font.color.rgb == "FFFF0000"
    assert ws["D2"].alignment.wrap_text and ws["D2"].border.left.style == "thin"
    assert ws.row_dimensions[2].height == 30 and ws.column_dimensions["D"].width == 50


def test_append_keeps_existing_and_repeated_ids(tmp_path):
    p = str(tmp_path / "day.xlsx")
    excel_io.create_new_workbook(p)
    rows = parse_file(SAMPLE).rows
    r1 = excel_io.append_rows(p, rows, backup=False)
    assert (r1["first_row"], r1["last_row"]) == (2, 11)
    r2 = excel_io.append_rows(p, rows[:2], backup=False)          # appends below, no dedupe
    assert r2["first_row"] == 12
    ws = load_workbook(p).active
    col_a = [ws.cell(r, 1).value for r in range(2, 14)]
    assert col_a.count("T20260101.0003") == 2 and col_a.count("T20260101.0001") == 3
    assert ws["A2"].value == "T20260101.0001"                    # first row untouched


def test_existing_rows_preserved_and_backup(tmp_path):
    p = str(tmp_path / "day.xlsx")
    excel_io.create_new_workbook(p)
    wb = load_workbook(p)
    wb.active["A2"], wb.active["E2"] = "T-OLD", "hand typed"
    wb.save(p)
    res = excel_io.append_rows(p, [TicketRow("T20260101.0099", "P1", "Assigned", "t", "d")])
    assert res["first_row"] == 3 and os.path.exists(res["backup"])
    ws = load_workbook(p).active
    assert ws["A2"].value == "T-OLD" and ws["E2"].value == "hand typed"


def test_text_starting_with_equals_is_not_a_formula(tmp_path):
    p = str(tmp_path / "f.xlsx")
    excel_io.create_new_workbook(p)
    excel_io.append_rows(p, [TicketRow("T1", "P1", "x", "t", "=SUM(A1:A2)")], backup=False)
    c = load_workbook(p).active["E2"]
    assert c.data_type == "s" and c.value == "=SUM(A1:A2)"


def test_count_existing_matches_is_informational(tmp_path):
    p = str(tmp_path / "d.xlsx")
    excel_io.create_new_workbook(p)
    rows = parse_file(SAMPLE).rows
    assert excel_io.count_existing_matches(p, rows) == 0
    excel_io.append_rows(p, rows, backup=False)
    assert excel_io.count_existing_matches(p, rows) == 10
